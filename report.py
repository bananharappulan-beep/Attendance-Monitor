"""Downloadable overall report: every branch, one section/sheet per branch.

Same columns and rules as the Daily Report on screen:
  LATE TIME       = punch in  - 09:10  (only when after 09:10)
  EARLY LEAVING   = 17:00     - punch out (only when before 17:00)
  WORKING HOURS   = punch out - punch in
  STATUS          = P (>= 7h) / H (4h to < 7h) / A (otherwise or no record)
"""
import io
import re
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import sheets_sync as S

LATE_AFTER = 9 * 60 + 10
EARLY_BEFORE = 17 * 60
HEAD = ["SL NO", "EMP NAME", "PUNCH IN TIME", "LATE TIME", "PUNCH OUT",
        "EARLY LEAVING TIME", "TOTAL WORKING HOURS", "STATUS"]
SUMMARY_HEAD = ["BRANCH", "EMPLOYEES", "PRESENT", "HALF-DAY", "ABSENT"]

# status colours (same as the web page): (background, text)
FILL = {"P": ("DCFCE7", "166534"), "H": ("FEF3C7", "92400E"), "A": ("FEE2E2", "991B1B")}


def _hm(m):
    return "" if m is None else f"{m // 60:02d}:{m % 60:02d}"


def _dmy(iso):
    return "-".join(reversed(iso.split("-")))


def branch_report(rows, iso):
    """-> {has_data, total, counts, rows:[[sl, name, in, late, out, early, work, st], ...]}"""
    by_key = {(r["d"], r["name"]): r for r in rows}
    names = sorted({r["name"] for r in rows}, key=str.casefold)
    has_data = any(r["d"] == iso for r in rows)
    counts = {"P": 0, "H": 0, "A": 0}
    out = []
    if has_data:
        for i, n in enumerate(names, 1):
            r = by_key.get((iso, n))
            in_m = S._minutes(r["inT"]) if r else None
            out_m = S._minutes(r["outT"]) if r else None
            work = out_m - in_m if in_m is not None and out_m is not None and out_m > in_m else 0
            st = S.status_for(r)
            counts[st] += 1
            late = _hm(in_m - LATE_AFTER) if in_m is not None and in_m > LATE_AFTER else ""
            early = _hm(EARLY_BEFORE - out_m) if out_m is not None and out_m < EARLY_BEFORE else ""
            out_txt = _hm(out_m) if out_m is not None else ("No out punch" if in_m is not None else "")
            out.append([i, n, _hm(in_m), late, out_txt, early, _hm(work), st])
    return {"has_data": has_data, "total": len(names), "counts": counts, "rows": out}


def collect(iso):
    data = S.get_all()
    return [(b, branch_report(data[b], iso)) for b in sorted(data)]


def _summary_rows(reports):
    rows = []
    tot = {"emp": 0, "P": 0, "H": 0, "A": 0}
    for b, rep in reports:
        if rep["has_data"]:
            c = rep["counts"]
            rows.append([b, rep["total"], c["P"], c["H"], c["A"]])
            tot["emp"] += rep["total"]
            for k in "PHA":
                tot[k] += c[k]
        else:
            rows.append([b, rep["total"], "No records for this date", "", ""])
    rows.append(["TOTAL", tot["emp"], tot["P"], tot["H"], tot["A"]])
    return rows


# =====================================================================
# Excel
# =====================================================================
def _safe_sheet_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "-", name).strip("'")[:31] or "Branch"
    base, n = t, 2
    while t.lower() in used:
        t = f"{base[:28]}-{n}"
        n += 1
    used.add(t.lower())
    return t


def _put(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        c.data_type = "s"          # never let a name be treated as a formula
    return c


def build_xlsx(iso):
    reports = collect(iso)
    wb = Workbook()
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="F1F5F9")
    bold = Font(bold=True)
    used = set()

    # ---- overall sheet ----
    ws = wb.active
    ws.title = _safe_sheet_title("Overall", used)
    ws["A1"] = f"Attendance Report - {_dmy(iso)} (all branches)"
    ws["A1"].font = Font(bold=True, size=14)
    for j, h in enumerate(SUMMARY_HEAD, 1):
        c = _put(ws, 3, j, h)
        c.font, c.fill, c.border = bold, head_fill, border
    rows = _summary_rows(reports)
    for i, row in enumerate(rows, 4):
        last = i == 3 + len(rows)
        for j, v in enumerate(row, 1):
            c = _put(ws, i, j, v)
            c.border = border
            if last:
                c.font = bold
                c.fill = head_fill
            if j > 1:
                c.alignment = Alignment(horizontal="center")
        if isinstance(row[2], str) and row[2].startswith("No records"):
            ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
            ws.cell(row=i, column=3).font = Font(italic=True, color="94A3B8")
    for col, w in zip("ABCDE", (28, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    # ---- one sheet per branch ----
    widths = (8, 30, 16, 12, 16, 20, 22, 10)
    for b, rep in reports:
        ws = wb.create_sheet(_safe_sheet_title(b, used))
        ws["A1"] = f"{b} - {_dmy(iso)}"
        ws["A1"].font = Font(bold=True, size=14)
        for j, h in enumerate(HEAD, 1):
            c = _put(ws, 3, j, h)
            c.font, c.fill, c.border = bold, head_fill, border
        if not rep["has_data"]:
            ws["A4"] = "No attendance records for this date."
            ws["A4"].font = Font(italic=True, color="94A3B8")
        for i, row in enumerate(rep["rows"], 4):
            for j, v in enumerate(row, 1):
                c = _put(ws, i, j, v)
                c.border = border
                if j == 4 or j == 6:
                    c.font = Font(color="DC2626")
                if j == 5 and v == "No out punch":
                    c.font = Font(italic=True, color="94A3B8")
                if j == 8:
                    bg, fg = FILL[v]
                    c.fill = PatternFill("solid", fgColor=bg)
                    c.font = Font(bold=True, color=fg)
                    c.alignment = Alignment(horizontal="center")
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# =====================================================================
# PDF
# =====================================================================
def _hex(h):
    return colors.HexColor("#" + h)


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(iso):
    reports = collect(iso)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title=f"Attendance Report {_dmy(iso)}")
    styles = getSampleStyleSheet()
    h1, h2 = styles["Heading1"], styles["Heading2"]
    grid, head_bg = colors.HexColor("#E2E8F0"), colors.HexColor("#F1F5F9")
    story = [Paragraph("Attendance Report - All Branches", h1),
             Paragraph(f"Date: {_dmy(iso)}", styles["Normal"]), Spacer(1, 8)]

    # ---- overall summary ----
    data = [SUMMARY_HEAD] + _summary_rows(reports)
    t = Table(data, colWidths=[200, 70, 70, 70, 70], repeatRows=1)
    style = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
             ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
             ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8),
             ("BACKGROUND", (0, 0), (-1, 0), head_bg),
             ("BACKGROUND", (0, -1), (-1, -1), head_bg),
             ("GRID", (0, 0), (-1, -1), 0.5, grid),
             ("ALIGN", (1, 0), (-1, -1), "CENTER")]
    for i, row in enumerate(data[1:], 1):
        if isinstance(row[2], str) and row[2].startswith("No records"):
            style += [("SPAN", (2, i), (4, i)),
                      ("TEXTCOLOR", (2, i), (4, i), colors.HexColor("#94A3B8"))]
    t.setStyle(TableStyle(style))
    story.append(t)

    # ---- one section per branch ----
    heads = ["SL NO", "EMP NAME", "PUNCH IN\nTIME", "LATE\nTIME", "PUNCH\nOUT",
             "EARLY LEAVING\nTIME", "TOTAL WORKING\nHOURS", "STATUS"]
    widths = [28, 140, 58, 42, 62, 72, 78, 40]
    for b, rep in reports:
        story += [PageBreak(), Paragraph(escape(b), h2),
                  Paragraph(f"Date: {_dmy(iso)}", styles["Normal"]), Spacer(1, 6)]
        if not rep["has_data"]:
            story.append(Paragraph("No attendance records for this date.", styles["Italic"]))
            continue
        c = rep["counts"]
        story += [Paragraph(f"Present: <b>{c['P']}</b> &nbsp;&nbsp; Half-day: <b>{c['H']}</b> "
                            f"&nbsp;&nbsp; Absent: <b>{c['A']}</b> &nbsp;&nbsp; "
                            f"Employees: <b>{rep['total']}</b>", styles["Normal"]), Spacer(1, 6)]
        t = Table([heads] + rep["rows"], colWidths=widths, repeatRows=1)
        style = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7),
                 ("FONT", (0, 1), (-1, -1), "Helvetica", 7.5),
                 ("BACKGROUND", (0, 0), (-1, 0), head_bg),
                 ("GRID", (0, 0), (-1, -1), 0.4, grid),
                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("ALIGN", (0, 0), (0, -1), "CENTER"),
                 ("ALIGN", (2, 0), (-1, -1), "CENTER"),
                 ("TEXTCOLOR", (3, 1), (3, -1), colors.HexColor("#DC2626")),
                 ("TEXTCOLOR", (5, 1), (5, -1), colors.HexColor("#DC2626"))]
        for i, row in enumerate(rep["rows"], 1):
            bg, fg = FILL[row[7]]
            style += [("BACKGROUND", (7, i), (7, i), _hex(bg)),
                      ("TEXTCOLOR", (7, i), (7, i), _hex(fg)),
                      ("FONT", (7, i), (7, i), "Helvetica-Bold", 7.5)]
            if row[4] == "No out punch":
                style += [("TEXTCOLOR", (4, i), (4, i), colors.HexColor("#94A3B8")),
                          ("FONT", (4, i), (4, i), "Helvetica-Oblique", 7)]
        t.setStyle(TableStyle(style))
        story.append(t)

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return buf
